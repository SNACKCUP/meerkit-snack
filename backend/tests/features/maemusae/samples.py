"""테스트 안에서 만드는 작은 샘플 파일 (PDF, DOCX, PPTX, XLSX). 저장소에 바이너리 파일을 두지 않으려고 코드로 만듭니다."""

import io
import zipfile


def pdf(pages: list[str]) -> bytes:
    """쪽마다 한 줄(ASCII)이 있는 PDF. 빈 문자열이면 글자 없는 쪽 (스캔 PDF 흉내)"""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    font = 3 + 2 * len(pages)
    for i, line in enumerate(pages):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font} 0 R >> >> /Contents {4 + 2 * i} 0 R >>"
        )
        stream = f"BT /F1 12 Tf 72 720 Td ({line}) Tj ET" if line else ""
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def docx(paragraphs: list[str]) -> bytes:
    """'# 제목' 으로 시작하는 문단은 제목1 스타일"""

    def para(text: str) -> str:
        if text.startswith("# "):
            return f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>{text[2:]}</w:t></w:r></w:p>'
        return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"

    ns = "http://schemas.openxmlformats.org"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr(
            "[Content_Types].xml",
            f'<?xml version="1.0"?><Types xmlns="{ns}/package/2006/content-types">'
            f'<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            f'<Default Extension="xml" ContentType="application/xml"/>'
            f'<Override PartName="/word/document.xml" ContentType="application/'
            f'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        zf.writestr(
            "_rels/.rels",
            f'<?xml version="1.0"?><Relationships xmlns="{ns}/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{ns}/officeDocument/2006/relationships/officeDocument" '
            f'Target="word/document.xml"/></Relationships>',
        )
        zf.writestr(
            "word/document.xml",
            f'<?xml version="1.0"?><w:document xmlns:w="{ns}/wordprocessingml/2006/main"><w:body>'
            + "".join(para(p) for p in paragraphs)
            + "</w:body></w:document>",
        )
    return buffer.getvalue()


def pptx(slides: list[tuple[str, str]]) -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    for title, body in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def xlsx(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
