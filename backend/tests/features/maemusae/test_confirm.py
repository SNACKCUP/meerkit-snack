"""매무새 요구사항 확정 (M4-1): 확정·확정 후 수정·확정한 문서 목록·Excel, 경고 개수, 권한, 로그, 업로드부터 Excel 까지 한 흐름"""

import io
import logging
from contextlib import nullcontext

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update

from app.features.maemusae.analysis import runner
from app.features.maemusae.analysis.models import Requirement, RequirementSource
from app.features.maemusae.documents.models import Document, SourcePage
from app.features.maemusae.projects.models import MEMBER, ProjectMember
from app.users.models import User

PROJECTS = "/api/mms/projects"
PROJECT = {"name": "쇼핑몰 리뉴얼", "start_date": "2026-11-01", "end_date": "2026-12-31"}


@pytest.fixture(autouse=True)
def _app_logs_visible():
    # app 로거는 propagate=False 라서 caplog 가 받도록 테스트 동안만 켭니다 (하네스 pytest-app-logging)
    logger = logging.getLogger("app")
    logger.propagate = True
    yield
    logger.propagate = False


@pytest.fixture(autouse=True)
def _runner_uses_test_session(db, monkeypatch):
    monkeypatch.setattr(runner, "open_session", lambda: nullcontext(db))


@pytest.fixture
def headers(db, make_user, login):
    def _headers(login_id: str, *, admin: bool = False) -> dict[str, str]:
        if db.scalar(select(User).where(User.login_id == login_id)) is None:
            make_user(login_id, admin=admin, name=f"{login_id} 이름")
        return login(login_id)

    return _headers


@pytest.fixture
def owner(headers):
    return headers("owner01")


@pytest.fixture
def docs_url(client, owner):
    project = client.post(PROJECTS, json=PROJECT, headers=owner).json()
    [document] = client.get(f"{PROJECTS}/{project['id']}/documents", headers=owner).json()
    return f"{PROJECTS}/{project['id']}/documents/{document['id']}"


def _doc_id(url: str) -> int:
    return int(url.rsplit("/", 1)[1])


def _project_url(url: str) -> str:
    return url.split("/documents/")[0]


@pytest.fixture
def member(db, headers, docs_url):
    hdrs = headers("member01")
    user = db.scalar(select(User).where(User.login_id == "member01"))
    db.add(ProjectMember(project_id=int(docs_url.split("/")[4]), user_id=user.id, role=MEMBER))
    db.commit()
    return hdrs


@pytest.fixture
def seeded(client, db, owner, docs_url):
    """분석이 끝난 문서 흉내: 직접 쓴 글 1개 + 후보 3개 (REQ-001~003, REQ-002 는 질문, REQ-003 은 AI 값 확인), READY"""
    assert client.post(f"{docs_url}/sources", data={"text": "회원은 이메일로 가입한다."}, headers=owner).status_code == 201
    page = db.scalar(select(SourcePage))
    doc_id = _doc_id(docs_url)
    rows = [
        Requirement(document_id=doc_id, seq=1, name="이메일 회원가입", kind="FUNCTIONAL", priority="HIGH",
                    description="이메일로 가입"),
        Requirement(document_id=doc_id, seq=2, name="카드 결제", kind="FUNCTIONAL", priority="MEDIUM",
                    description="=HYPERLINK(\"http://x\")", ai_question="간편결제도 되나요?"),
        Requirement(document_id=doc_id, seq=3, name="응답 3초", kind="NON_FUNCTIONAL", priority="LOW",
                    priority_defaulted=True),
    ]
    db.add_all(rows)
    db.flush()
    db.add_all([
        RequirementSource(requirement_id=rows[0].id, source_page_id=page.id, quote="이메일로 가입한다", verified=True,
                          sort_order=1),
        RequirementSource(requirement_id=rows[1].id, source_page_id=page.id, quote="지어낸 문장", verified=False,
                          sort_order=1),
    ])
    db.execute(update(Document).where(Document.id == doc_id).values(status="READY"))
    db.commit()
    return {r.seq: r.id for r in rows}


def _document(client, url, hdrs) -> dict:
    response = client.get(url, headers=hdrs)
    assert response.status_code == 200, response.text
    return response.json()


def _uncheck(client, url, hdrs, requirement_id, version=1):
    response = client.patch(f"{url}/requirements/{requirement_id}", json={"version": version, "is_included": False}, headers=hdrs)
    assert response.status_code == 200, response.text


def _sheet(content: bytes):
    return load_workbook(io.BytesIO(content)).active


def _table(sheet) -> list[list]:
    rows = []
    for row in sheet.iter_rows(min_row=10, values_only=True):
        if row[0] is None or not str(row[0]).startswith("REQ-"):
            break
        rows.append(list(row))
    return rows


# ---- 경고 개수 ----


def test_document_has_warning_counts_of_checked_candidates(client, owner, docs_url, seeded):
    document = _document(client, docs_url, owner)
    assert (document["open_question_count"], document["defaulted_count"]) == (1, 1)

    client.put(f"{docs_url}/requirements/{seeded[2]}/notes/ANSWER", json={"text": "네"}, headers=owner)
    _uncheck(client, docs_url, owner, seeded[3])
    document = _document(client, docs_url, owner)
    assert (document["open_question_count"], document["defaulted_count"]) == (0, 0)  # 답했고, 체크 해제는 안 셈


# ---- 확정 ----


def test_confirm_keeps_checked_only_and_locks_document(client, db, owner, docs_url, seeded):
    _uncheck(client, docs_url, owner, seeded[1])
    response = client.post(f"{docs_url}/confirm", headers=owner)
    assert response.status_code == 201, response.text
    result = response.json()
    assert (result["revision"], result["item_count"]) == (1, 2)

    document = _document(client, docs_url, owner)
    assert (document["status"], document["current_revision"], document["confirmed_by_name"]) == ("CONFIRMED", 1, "owner01 이름")
    assert document["confirmed_at"] is not None

    detail = client.get(f"{docs_url}/versions/{result['version_id']}", headers=owner).json()
    assert (detail["kind"], detail["revision"], detail["version_no"]) == ("CONFIRMED", 1, None)
    assert [i["code"] for i in detail["items"]] == ["REQ-002", "REQ-003"]
    assert detail["items"][0]["sources"][0]["quote"] == "지어낸 문장"

    # 확정 기록은 그 순간 그대로, AI 값 확인 표시는 풀림 (version 은 그대로)
    third = db.get(Requirement, seeded[3])
    db.refresh(third)
    assert (third.priority_defaulted, third.version) == (False, 1)

    # 확정된 문서는 고치기·분석·버전 저장이 막힘
    calls = [
        client.patch(f"{docs_url}/requirements/{seeded[2]}", json={"version": 1, "name": "x"}, headers=owner),
        client.post(f"{docs_url}/versions", headers=owner),
        client.post(f"{docs_url}/analyze", json={"restart": True}, headers=owner),
        client.post(f"{docs_url}/confirm", headers=owner),
    ]
    assert [(c.status_code, c.json()["code"]) for c in calls] == [(409, "DOCUMENT_CONFIRMED")] * 4


def test_reopen_then_confirm_again_makes_revision_two(client, db, owner, docs_url, seeded):
    first = client.post(f"{docs_url}/confirm", headers=owner).json()
    assert client.post(f"{docs_url}/reopen", headers=owner).status_code == 204
    document = _document(client, docs_url, owner)
    assert (document["status"], document["current_revision"], document["changed_since_save"]) == ("READY", 1, False)

    _uncheck(client, docs_url, owner, seeded[3])  # 다시 고칠 수 있음
    second = client.post(f"{docs_url}/confirm", headers=owner).json()
    assert (second["revision"], second["item_count"]) == (2, 2)

    listed = client.get(f"{_project_url(docs_url)}/confirmed", headers=owner).json()
    assert [(d["version_id"], d["revision"], d["item_count"], d["kind_label"]) for d in listed] == [
        (second["version_id"], 2, 2, "요구사항정의서"),
        (first["version_id"], 1, 3, "요구사항정의서"),
    ]
    assert listed[0]["confirmed_by_name"] == "owner01 이름"
    # 첫 확정 기록은 바뀌지 않음
    old = client.get(f"{docs_url}/versions/{first['version_id']}", headers=owner).json()
    assert old["item_count"] == 3

    # 버전 기록 창에도 확정이 같이 보임
    kinds = [(v["kind"], v["revision"]) for v in client.get(f"{docs_url}/versions", headers=owner).json()]
    assert kinds == [("CONFIRMED", 2), ("CONFIRMED", 1)]

    assert client.post(f"{docs_url}/reopen", headers=owner).status_code == 204
    assert client.post(f"{docs_url}/reopen", headers=owner).json()["code"] == "DOCUMENT_NOT_CONFIRMED"


@pytest.mark.parametrize(
    ("status", "http", "code"),
    [("EMPTY", 409, "DOCUMENT_NOT_READY"), ("FAILED", 409, "DOCUMENT_NOT_READY"), ("ANALYZING", 409, "DOCUMENT_BUSY")],
)
def test_confirm_needs_ready_document(client, db, owner, docs_url, seeded, status, http, code):
    db.execute(update(Document).where(Document.id == _doc_id(docs_url)).values(status=status))
    db.commit()
    response = client.post(f"{docs_url}/confirm", headers=owner)
    assert (response.status_code, response.json()["code"]) == (http, code)


def test_confirm_needs_a_checked_candidate(client, owner, docs_url, seeded):
    for seq in (1, 2, 3):
        _uncheck(client, docs_url, owner, seeded[seq])
    response = client.post(f"{docs_url}/confirm", headers=owner)
    assert (response.status_code, response.json()["code"]) == (400, "NOTHING_TO_CONFIRM")
    assert _document(client, docs_url, owner)["status"] == "READY"


# ---- 권한 ----


def test_only_project_admin_confirms_and_reopens(client, headers, owner, member, docs_url, seeded):
    master = headers("master01", admin=True)
    outsider = headers("outsider01")

    response = client.post(f"{docs_url}/confirm", headers=member)
    assert (response.status_code, response.json()["code"]) == (403, "PROJECT_ADMIN_REQUIRED")
    response = client.post(f"{docs_url}/confirm", headers=master)
    assert (response.status_code, response.json()["code"]) == (403, "PROJECT_MEMBER_REQUIRED")
    assert client.post(f"{docs_url}/confirm", headers=outsider).status_code == 404

    version_id = client.post(f"{docs_url}/confirm", headers=owner).json()["version_id"]
    assert client.post(f"{docs_url}/reopen", headers=member).json()["code"] == "PROJECT_ADMIN_REQUIRED"
    assert client.post(f"{docs_url}/reopen", headers=master).json()["code"] == "PROJECT_MEMBER_REQUIRED"

    # 목록은 MASTER 도 보지만, Excel 은 멤버만 (고객 문서를 밖으로 가져가지 않음)
    assert client.get(f"{_project_url(docs_url)}/confirmed", headers=master).status_code == 200
    assert client.get(f"{_project_url(docs_url)}/confirmed", headers=outsider).status_code == 404
    excel = f"{docs_url}/versions/{version_id}/excel"
    assert client.get(excel, headers=member).status_code == 200
    response = client.get(excel, headers=master)
    assert (response.status_code, response.json()["code"]) == (403, "PROJECT_MEMBER_REQUIRED")
    assert client.get(excel, headers=outsider).status_code == 404


# ---- Excel ----


def test_excel_is_built_from_confirmed_record(client, owner, docs_url, seeded):
    _uncheck(client, docs_url, owner, seeded[1])
    version_id = client.post(f"{docs_url}/confirm", headers=owner).json()["version_id"]
    response = client.get(f"{docs_url}/versions/{version_id}/excel", headers=owner)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    assert "filename*=UTF-8''%EC%87%BC%ED%95%91%EB%AA%B0" in disposition  # 쇼핑몰
    assert disposition.endswith("%ED%99%95%EC%A0%951.xlsx")  # 확정1
    assert response.headers["cache-control"] == "private, no-store"

    sheet = _sheet(response.content)
    assert sheet.title == "요구사항정의서"
    assert sheet["A1"].value == "요구사항정의서"
    head = {sheet.cell(r, 1).value: sheet.cell(r, 2).value for r in range(2, 8)}
    assert head["프로젝트명"] == "쇼핑몰 리뉴얼"
    assert head["확정 번호"] == "확정 1 (확정)"
    assert head["확정한 사람"] == "owner01 이름"
    assert head["요구사항 수"] == "2개"
    assert head["확정일"].endswith("(한국 시간)")
    assert [c.value for c in sheet[9]] == ["ID", "대분류", "중분류", "소분류", "요구사항명", "구분", "우선순위", "설명", "근거"]
    assert sheet.freeze_panes == "A10"

    table = _table(sheet)
    assert [r[0] for r in table] == ["REQ-002", "REQ-003"]
    second = table[0]
    assert (second[4], second[5], second[6]) == ("카드 결제", "기능", "보통")  # 코드 대신 공통코드 이름
    assert second[7] == "=HYPERLINK(\"http://x\")"  # 글자 그대로, 수식이 아님
    assert sheet.cell(10, 8).data_type == "s"
    assert second[8] == "직접 입력 · 직접 입력 (위치 확인 필요)\n\"지어낸 문장\""
    # REQ-001 을 뺐으므로 번호가 빈다는 안내
    assert "빠진 번호" in (sheet.cell(10 + len(table) + 1, 1).value or "")

    # 확정 뒤 후보를 고쳐도 확정 기록의 Excel 은 그대로
    client.post(f"{docs_url}/reopen", headers=owner)
    client.patch(f"{docs_url}/requirements/{seeded[2]}", json={"version": 1, "name": "바뀐 이름"}, headers=owner)
    again = _table(_sheet(client.get(f"{docs_url}/versions/{version_id}/excel", headers=owner).content))
    assert again[0][4] == "카드 결제"


def test_excel_sorts_and_merges_by_category():
    from datetime import UTC, datetime

    from app.features.maemusae.confirm.excel import SheetMeta, build_requirements_xlsx

    def item(code, major=None, middle=None, minor=None):
        return {"code": code, "name": code, "kind": "FUNCTIONAL", "priority": "HIGH", "description": "",
                "category_major": major, "category_middle": middle, "category_minor": minor, "sources": []}

    items = [
        item("REQ-001"),
        item("REQ-002", "회원 관리", "로그인", "소셜 로그인"),
        item("REQ-003", "주문", "결제", "카드"),
        item("REQ-004", "회원 관리", "로그인", "아이디 로그인"),
        item("REQ-005", "회원 관리", "가입"),
    ]
    meta = SheetMeta("p", "요구사항정의서", 2, datetime(2026, 10, 31, 7, 32, tzinfo=UTC), "김다나")
    sheet = _sheet(build_requirements_xlsx(meta, items, str, str))
    table = _table(sheet)
    assert [r[0] for r in table] == ["REQ-003", "REQ-005", "REQ-002", "REQ-004", "REQ-001"]  # 분류 순(가나다), 없는 것은 맨 아래
    assert sheet.cell(4, 2).value == "확정 2 (수정 후 확정)"
    assert sheet.cell(5, 2).value == "2026-10-31 16:32 (한국 시간)"
    merged = {str(r) for r in sheet.merged_cells.ranges}
    assert "B11:B13" in merged  # 회원 관리 3줄
    assert "C12:C13" in merged  # 회원 관리 > 로그인 2줄
    assert not any(r.startswith("B14") for r in merged)  # 빈 분류는 병합 안 함
    assert sheet.cell(10 + len(table) + 1, 1).value is None  # 번호가 비지 않으면 안내 없음


def test_excel_of_saved_version_or_other_document_is_404(client, owner, docs_url, seeded):
    saved = client.post(f"{docs_url}/versions", headers=owner).json()
    response = client.get(f"{docs_url}/versions/{saved['id']}/excel", headers=owner)
    assert (response.status_code, response.json()["code"]) == (404, "CONFIRMED_VERSION_NOT_FOUND")

    confirmed = client.post(f"{docs_url}/confirm", headers=owner).json()
    other = client.post(PROJECTS, json=PROJECT, headers=owner).json()
    [other_doc] = client.get(f"{PROJECTS}/{other['id']}/documents", headers=owner).json()
    url = f"{PROJECTS}/{other['id']}/documents/{other_doc['id']}/versions/{confirmed['version_id']}/excel"
    assert client.get(url, headers=owner).json()["code"] == "CONFIRMED_VERSION_NOT_FOUND"


def test_unsafe_project_name_is_cleaned_in_file_name(client, owner, docs_url, seeded):
    project_url = _project_url(docs_url)
    client.patch(project_url, json={"name": "A/B:C"}, headers=owner)
    version_id = client.post(f"{docs_url}/confirm", headers=owner).json()["version_id"]
    disposition = client.get(f"{docs_url}/versions/{version_id}/excel", headers=owner).headers["content-disposition"]
    assert "filename*=UTF-8''A_B_C_" in disposition


# ---- 로그 ----


def test_logs_have_ids_not_content(client, owner, docs_url, seeded, caplog):
    with caplog.at_level(logging.INFO, logger="app.features.maemusae"):
        version_id = client.post(f"{docs_url}/confirm", headers=owner).json()["version_id"]
        client.get(f"{docs_url}/versions/{version_id}/excel", headers=owner)
        client.post(f"{docs_url}/reopen", headers=owner)
    assert "요구사항 확정 document_id=" in caplog.text and "revision=1 items=3" in caplog.text
    assert "요구사항정의서 Excel" in caplog.text and "확정 후 수정" in caplog.text
    for secret in ("이메일 회원가입", "카드 결제", "지어낸 문장", "쇼핑몰"):
        assert secret not in caplog.text


# ---- 한 흐름 (08 계획: 업로드부터 확정·Excel 까지 한 테스트로) ----


def test_flow_upload_analyze_review_confirm_excel(client, db, owner, docs_url, fake_llm):
    assert client.post(f"{docs_url}/sources", data={"text": "회원은 이메일로 가입한다. 결제는 카드로 한다."}, headers=owner).status_code == 201
    fake_llm.responses = [{
        "requirements": [
            {"name": "이메일 회원가입", "description": "이메일로 가입", "kind": "FUNCTIONAL", "priority": "HIGH",
             "question": None, "evidence": [{"page": "P1", "quote": "이메일로 가입한다"}]},
            {"name": "카드 결제", "description": "카드 결제", "kind": "FUNCTIONAL", "priority": "MEDIUM",
             "question": "간편결제도 되나요?", "evidence": [{"page": "P1", "quote": "카드로 한다"}]},
        ]
    }]
    assert client.post(f"{docs_url}/analyze", json={}, headers=owner).status_code == 202
    assert _document(client, docs_url, owner)["status"] == "READY"

    items = client.get(f"{docs_url}/requirements", headers=owner).json()["items"]
    assert [i["code"] for i in items] == ["REQ-001", "REQ-002"]
    second = items[1]
    client.put(f"{docs_url}/requirements/{second['id']}/notes/ANSWER", json={"text": "아니요"}, headers=owner)
    client.patch(f"{docs_url}/requirements/{second['id']}", json={"version": second["version"], "priority": "HIGH"}, headers=owner)
    assert _document(client, docs_url, owner)["open_question_count"] == 0

    confirmed = client.post(f"{docs_url}/confirm", headers=owner).json()
    table = _table(_sheet(client.get(f"{docs_url}/versions/{confirmed['version_id']}/excel", headers=owner).content))
    assert [(r[0], r[4], r[6]) for r in table] == [("REQ-001", "이메일 회원가입", "높음"), ("REQ-002", "카드 결제", "높음")]
    assert table[0][8] == "직접 입력 · 직접 입력\n\"이메일로 가입한다\""

    assert client.post(f"{docs_url}/reopen", headers=owner).status_code == 204
    assert client.post(f"{docs_url}/confirm", headers=owner).json()["revision"] == 2
