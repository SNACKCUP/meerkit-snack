"""확정한 요구사항정의서를 Excel(.xlsx)로 그립니다 (계획 19 결정 B1: 코드로 그리는 고정 양식).

    1행      요구사항정의서 (제목)
    2~7행    머리 칸: 프로젝트명, 문서, 확정 번호, 확정일, 확정한 사람, 요구사항 수
    9행      표 머리줄 (고정): ID, 대분류, 중분류, 소분류, 요구사항명, 구분, 우선순위, 설명, 근거
    10행~    후보. 분류(대→중→소) 순, 그 안은 REQ 번호 순. 분류 없는 후보는 맨 아래. 같은 대·중분류 칸은 병합
    마지막    번호가 비면 안내 한 줄 (결정 D1: 번호는 바꾸지 않음)

- 분류 칸은 후보에 대·중·소분류가 생기면(M4-2, 10/29 결정 F) 채워지고, 그 전 기록은 빈칸
- 고객 문서 내용이 들어 있어 파일은 저장하지 않고 요청마다 메모리에서 만듭니다
- 모든 글은 문자열 칸으로 넣어서 '=' 로 시작해도 수식으로 저장되지 않습니다 (수식 주입 방지)
"""

import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

KST = timezone(timedelta(hours=9), "KST")  # 한국은 서머타임이 없어 고정 +9

COLUMNS = [  # (머리줄, 너비)
    ("ID", 10),
    ("대분류", 14),
    ("중분류", 14),
    ("소분류", 14),
    ("요구사항명", 32),
    ("구분", 10),
    ("우선순위", 9),
    ("설명", 48),
    ("근거", 56),
]
TABLE_HEADER_ROW = 9
GAP_NOTE = "빠진 번호는 체크를 해제했거나 지운 후보예요."

_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="F2F1EC")
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_CODE = re.compile(r"(\d+)$")


@dataclass(frozen=True)
class SheetMeta:
    project_name: str
    title: str
    revision: int
    confirmed_at: datetime
    confirmed_by: str


def _text(cell: Cell, value: str | int | None) -> Cell:
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # '=…' 도 글자 그대로 (openpyxl 은 '=' 로 시작하면 수식으로 저장함)
    return cell


def _seq(code: str) -> int:
    match = _CODE.search(code)
    return int(match.group(1)) if match else 0


def _category(item: dict, level: str) -> str:
    return (item.get(f"category_{level}") or "").strip()


def _sort_key(item: dict) -> tuple:
    major, middle, minor = (_category(item, lv) for lv in ("major", "middle", "minor"))
    return (not major, major, middle, minor, _seq(item["code"]))


def _evidence(item: dict) -> str:
    lines = []
    for s in item.get("sources", []):
        where = s["location_label"] + (f"~{s['end_location_label']}" if s.get("end_location_label") else "")
        mark = "" if s["verified"] else " (위치 확인 필요)"
        lines.append(f"{s['filename']} · {where}{mark}\n\"{s['quote']}\"")
    return "\n".join(lines)


def _merge_runs(sheet: Worksheet, column: int, keys: list[tuple], first_row: int) -> None:
    """같은 값이 이어지는 칸을 위아래로 병합 (빈 값은 병합하지 않음)"""
    start = 0
    for i in range(1, len(keys) + 1):
        if i < len(keys) and keys[i] == keys[start]:
            continue
        if i - start > 1 and keys[start][-1]:
            sheet.merge_cells(start_row=first_row + start, end_row=first_row + i - 1, start_column=column, end_column=column)
        start = i


def build_requirements_xlsx(
    meta: SheetMeta,
    items: list[dict],
    kind_label: Callable[[str], str],
    priority_label: Callable[[str], str],
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "요구사항정의서"
    last_col = get_column_letter(len(COLUMNS))

    title = _text(sheet["A1"], meta.title)
    title.font = Font(size=16, bold=True)
    sheet.merge_cells(f"A1:{last_col}1")

    kind = "확정" if meta.revision == 1 else "수정 후 확정"
    head = [
        ("프로젝트명", meta.project_name),
        ("문서", meta.title),
        ("확정 번호", f"확정 {meta.revision} ({kind})"),
        ("확정일", meta.confirmed_at.astimezone(KST).strftime("%Y-%m-%d %H:%M (한국 시간)")),
        ("확정한 사람", meta.confirmed_by),
        ("요구사항 수", f"{len(items)}개"),
    ]
    for offset, (label, value) in enumerate(head):
        row = 2 + offset
        _text(sheet.cell(row, 1, None), label).font = Font(bold=True)
        _text(sheet.cell(row, 2, None), value)
        sheet.merge_cells(start_row=row, end_row=row, start_column=2, end_column=len(COLUMNS))

    for col, (name, width) in enumerate(COLUMNS, 1):
        cell = _text(sheet.cell(TABLE_HEADER_ROW, col, None), name)
        cell.font = Font(bold=True)
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = sheet.cell(TABLE_HEADER_ROW + 1, 1)

    rows = sorted(items, key=_sort_key)
    first = TABLE_HEADER_ROW + 1
    for offset, item in enumerate(rows):
        values = [
            item["code"],
            _category(item, "major"),
            _category(item, "middle"),
            _category(item, "minor"),
            item["name"],
            kind_label(item["kind"]),
            priority_label(item["priority"]),
            item.get("description") or "",
            _evidence(item),
        ]
        for col, value in enumerate(values, 1):
            cell = _text(sheet.cell(first + offset, col, None), value)
            cell.border = _BORDER
            cell.alignment = _WRAP_TOP
    majors = [(_category(i, "major"),) for i in rows]
    middles = [(_category(i, "major"), _category(i, "middle")) for i in rows]
    _merge_runs(sheet, 2, majors, first)
    _merge_runs(sheet, 3, middles, first)

    seqs = [_seq(i["code"]) for i in items]
    if seqs and max(seqs) > len(seqs):
        note = _text(sheet.cell(first + len(rows) + 1, 1, None), GAP_NOTE)
        note.font = Font(italic=True, color="6C6A62")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
